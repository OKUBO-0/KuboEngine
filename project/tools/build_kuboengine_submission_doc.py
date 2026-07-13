from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "output" / "documents"
DOCX_PATH = OUT_DIR / "KuboEngine_program_explanation_A4.docx"

FONT_BODY = "Noto Sans JP"
FONT_HEADING = "Noto Sans JP Black"

NAVY = "15324A"
BLUE = "27678A"
GREEN = "39735D"
ORANGE = "B46A32"
PALE_BLUE = "EAF2F6"
PALE_GREEN = "EAF3EE"
PALE_ORANGE = "F8EFE7"
PALE_GRAY = "F3F5F6"
GRAY = "5D6870"
WHITE = "FFFFFF"
INK = "1C252B"
BORDER = "202020"

TABLE_WIDTH = 9800


def _set_font(run, size=9.3, bold=False, color=INK, font=None):
    font = font or (FONT_HEADING if bold else FONT_BODY)
    run.font.name = font
    r_pr = run._element.get_or_add_rPr()
    r_pr.rFonts.set(qn("w:ascii"), font)
    r_pr.rFonts.set(qn("w:hAnsi"), font)
    r_pr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def _p_style(p, before=0, after=4, line=1.12, keep=False):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = keep


def add_text(p, text, size=9.3, bold=False, color=INK, font=None):
    run = p.add_run(text)
    _set_font(run, size=size, bold=bold, color=color, font=font)
    return run


def add_para(doc, text, size=9.3, color=INK, bold=False, after=4, line=1.12, align=None):
    p = doc.add_paragraph()
    _p_style(p, after=after, line=line)
    if align is not None:
        p.alignment = align
    add_text(p, text, size=size, color=color, bold=bold)
    return p


def add_heading(doc, text, level=1, color=NAVY):
    p = doc.add_paragraph(style=f"Heading {level}")
    _p_style(p, before=7 if level == 1 else 4, after=4, line=1.0, keep=True)
    add_text(p, text, size=14.2 if level == 1 else 11.2, bold=True, color=color)
    return p


def add_kicker(doc, text, color=BLUE):
    p = doc.add_paragraph()
    _p_style(p, after=2, line=1.0, keep=True)
    add_text(p, text, size=8.1, bold=True, color=color)


def add_bullet(doc, text, color=INK):
    p = doc.add_paragraph(style="List Bullet")
    _p_style(p, after=2.4, line=1.08)
    add_text(p, text, size=8.8, color=color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders_for_cell(cell, color=BORDER, size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    b = tc_pr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders")
        tc_pr.append(b)
    for edge in ("top", "left", "start", "bottom", "right", "end"):
        node = b.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            b.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def margins_for_cell(cell, top=95, start=120, bottom=95, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def borders_for_table(table, color=BORDER, size="8"):
    tbl_pr = table._tbl.tblPr
    b = tbl_pr.find(qn("w:tblBorders"))
    if b is None:
        b = OxmlElement("w:tblBorders")
        tbl_pr.append(b)
    for edge in ("top", "left", "start", "bottom", "right", "end", "insideH", "insideV"):
        node = b.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            b.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")

    borders_for_table(table)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            margins_for_cell(cell)
            borders_for_cell(cell)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_table(doc, headers, rows, widths, header_fill=PALE_BLUE, font_size=8.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        shade(cell, header_fill)
        p = cell.paragraphs[0]
        _p_style(p, after=0, line=1.04)
        add_text(p, header, size=font_size, bold=True, color=NAVY)
    repeat_header(table.rows[0])

    for row in rows:
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            p = cells[idx].paragraphs[0]
            _p_style(p, after=0, line=1.08)
            add_text(p, str(text), size=font_size, color=INK, bold=(idx == 0 and len(headers) == 2))

    set_table_geometry(table, widths)
    return table


def add_flow(doc, items, fill=PALE_BLUE, accent=BLUE, font_size=8.0):
    table = doc.add_table(rows=1, cols=len(items))
    table.style = "Table Grid"
    widths = [TABLE_WIDTH // len(items)] * len(items)
    widths[-1] += TABLE_WIDTH - sum(widths)
    for idx, item in enumerate(items):
        cell = table.cell(0, idx)
        shade(cell, fill if idx % 2 == 0 else WHITE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _p_style(p, after=0, line=1.06)
        add_text(p, item, size=font_size, bold=True, color=accent)
    set_table_geometry(table, widths)
    return table


def add_callout(doc, label, text, fill=PALE_GRAY, accent=BLUE, label_w=1550, font_size=8.4):
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    shade(table.cell(0, 0), accent)
    shade(table.cell(0, 1), fill)

    p0 = table.cell(0, 0).paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _p_style(p0, after=0, line=1.0)
    add_text(p0, label, size=8.1, bold=True, color=WHITE)

    p1 = table.cell(0, 1).paragraphs[0]
    _p_style(p1, after=0, line=1.08)
    add_text(p1, text, size=font_size, color=INK)
    set_table_geometry(table, [label_w, TABLE_WIDTH - label_w])
    return table


def add_codebox(doc, label, text, accent=NAVY, fill="FAFBFC", font_size=7.6):
    table = doc.add_table(rows=2, cols=1)
    table.style = "Table Grid"
    shade(table.cell(0, 0), accent)
    p0 = table.cell(0, 0).paragraphs[0]
    _p_style(p0, after=0, line=1.0)
    add_text(p0, label, size=8.0, bold=True, color=WHITE)
    shade(table.cell(1, 0), fill)
    p1 = table.cell(1, 0).paragraphs[0]
    _p_style(p1, after=0, line=1.03)
    add_text(p1, text, size=font_size, color=INK, font=FONT_BODY)
    set_table_geometry(table, [TABLE_WIDTH])
    return table


def add_spacer(doc, points=4):
    p = doc.add_paragraph()
    _p_style(p, after=points, line=1.0)


def page_break(doc):
    doc.add_page_break()


def add_page_header(section):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _p_style(p, after=0, line=1.0)
    add_text(p, "KuboEngine / DirectXGame  |  PROGRAM EXPLANATION", size=7.4, bold=True, color=GRAY)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _p_style(fp, after=0, line=1.0)
    add_text(fp, "KuboEngine  |  ", size=7.3, color=GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    fp._p.append(fld)


def configure_doc(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.42)
    section.bottom_margin = Cm(1.30)
    section.left_margin = Cm(1.60)
    section.right_margin = Cm(1.60)
    section.header_distance = Cm(0.65)
    section.footer_distance = Cm(0.60)
    add_page_header(section)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_BODY
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
    normal.font.size = Pt(9.3)
    normal.font.color.rgb = RGBColor.from_string(INK)

    for name in ("Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = FONT_HEADING
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_HEADING)
        style.font.color.rgb = RGBColor.from_string(NAVY)

    lb = doc.styles["List Bullet"]
    lb.font.name = FONT_BODY
    lb._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
    lb.paragraph_format.left_indent = Cm(0.62)
    lb.paragraph_format.first_line_indent = Cm(-0.30)


def build():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_doc(doc)

    # Page 1
    add_kicker(doc, "PROGRAM EXPLANATION  |  WINDOWS / DIRECTX 12", ORANGE)
    p = doc.add_paragraph()
    _p_style(p, after=3, line=0.95, keep=True)
    add_text(p, "KuboEngine", size=24.0, bold=True, color=NAVY)
    p = doc.add_paragraph()
    _p_style(p, after=7, line=1.0, keep=True)
    add_text(p, "C++／DirectX 12自作エンジン上で制作した3Dサバイバルアクション", size=12.2, bold=True, color=BLUE)

    add_para(
        doc,
        "C++20とDirectX 12でゲームエンジンを構築し、その上でTitle、Play、Resultまでを持つゲームを制作しました。プレイヤーは敵を倒して経験値を集め、レベルアップ時に武器や能力を選びながらボス戦まで生き残ります。獲得コインはResultで集計し、Titleの恒久強化やキャラクター解放へ引き継ぎます。",
        size=9.4,
        after=6,
        line=1.16,
    )
    add_flow(doc, ["C++20", "DirectX 12", "HLSL", "Dear ImGui", "Assimp", "CSV"], PALE_BLUE, BLUE, 8.0)
    add_spacer(doc, 4)
    add_table(
        doc,
        ["範囲", "実装内容"],
        [
            ("エンジン", "DirectX 12基盤、描画、入力、音声、モデル、スプライト、パーティクル、カメラ、シーン管理"),
            ("ゲーム", "Player、Enemy、Weapon、HUD、LevelUp、GameSession、ボス戦、デバッグ表示"),
            ("検証", "CPU回帰テスト、シーン遷移ストレス、SRVテレメトリ、衝突CPU計測"),
        ],
        [2200, 7600],
        PALE_ORANGE,
        8.15,
    )
    add_heading(doc, "技術的な軸", 1, NAVY)
    add_callout(doc, "設計", "PlaySceneへ処理が集中しないよう、ゲーム進行、武器、敵、HUDを責務ごとに分けました。結果として、武器追加や進行変更の確認範囲を絞れます。", PALE_BLUE, BLUE)
    add_callout(doc, "描画", "CPUとGPUの参照タイミングのずれに対応するため、FrameContext、Fence、SRV再利用、ResourceBarrierで寿命と状態を追跡します。", PALE_GREEN, GREEN)
    add_callout(doc, "負荷対策", "全敵へ毎回OBB判定を行うのではなく、固定グリッドで候補を絞ってから詳細判定します。実測では84体時のCPU時間中央値が0.9304 msから0.1654 msへ下がりました。", PALE_ORANGE, ORANGE)

    # Page 2
    page_break(doc)
    add_kicker(doc, "01  ARCHITECTURE", BLUE)
    add_heading(doc, "エンジン層とゲーム層の分離", 1, NAVY)
    add_para(
        doc,
        "共通機能はengine/、作品固有の処理はgame/directxgame/に分けています。この境界を置いた理由は、ゲーム固有の変更が描画・入力・音声などの共通基盤へ混ざると、別作品へ使い回すときに切り離しにくくなるためです。",
        after=5,
    )
    add_flow(doc, ["main.cpp", "Game", "GameSceneFactory", "Title / Play / Result"], PALE_BLUE, BLUE, 8.0)
    add_spacer(doc, 4)
    add_codebox(
        doc,
        "簡略クラス図",
        "main.cpp -> Game -> GameSceneFactory -> TitleScene / PlayScene / ResultScene\n"
        "PlayScene -> GameplayFlowController / PlayerManager / PlayerWeaponController / EnemyManager / GameplayHudPresentation\n"
        "GameSession -> Title、Play、Resultをまたぐ結果・コイン・恒久強化を保持",
        BLUE,
        "FAFBFC",
        7.7,
    )
    add_heading(doc, "PlaySceneは構成の起点に絞る", 2, BLUE)
    add_para(
        doc,
        "PlaySceneへ更新、UI、演出を直接追加し続けると、1機能の変更でもシーン全体を確認する必要があります。そこでPlaySceneは各機能を保持して接続する役割に寄せ、処理は変更理由ごとのクラスへ分けました。",
        after=4,
    )
    add_table(
        doc,
        ["クラス", "主な役割"],
        [
            ("GameSceneFactory", "Title、Play、Resultシーンの生成"),
            ("PlayScene", "ゲームプレイを構成する機能の保持と接続"),
            ("GameplayFlowController", "Start、Playing、Boss、Pause、LevelUp、Deadの状態管理"),
            ("PlayerManager", "HP、経験値、レベル、能力値、武器強化の窓口"),
            ("EnemyManager", "敵の生成、EXP、衝突、ボス進行の管理"),
            ("GameSession", "結果、コイン、恒久強化、キャラクター情報の管理"),
        ],
        [2850, 6950],
        PALE_BLUE,
        8.0,
    )
    add_heading(doc, "ゲーム進行とシーン間データ", 2, GREEN)
    add_para(
        doc,
        "通常更新とポーズ、レベルアップ、死亡処理が同時に動かないよう、GameplayFlowControllerで状態を分けました。シーンをまたぐ値はGameSessionに集約し、Playは結果を記録し、Resultは表示し、Titleは強化へ使う、という役割にしています。",
        after=3,
    )
    add_flow(doc, ["Play\n結果を記録", "Result\n集計・表示", "Title\n強化・解放", "Play\n次のラン"], PALE_GREEN, GREEN, 8.0)

    # Page 3
    page_break(doc)
    add_kicker(doc, "02  DIRECTX 12 / RENDERING", GREEN)
    add_heading(doc, "DirectX 12のフレーム・リソース管理", 1, NAVY)
    add_para(
        doc,
        "DirectX 12では、CPUが次のフレームの準備を始めても、GPUが前のフレームのリソースをまだ参照している場合があります。CPU側の都合だけでCommandAllocatorや一時リソースを再利用すると、GPUが参照中の内容を上書きしてしまいます。",
        after=4,
    )
    add_codebox(
        doc,
        "CPU / GPU タイムライン",
        "時間 ->\n"
        "CPU  Frame 0を記録      Frame 1を記録      Frame 0を再利用\n"
        "     Allocator 0        Allocator 1        Fence 0完了後\n"
        "     Upload Arena 0     Upload Arena 1\n"
        "          | Fence 0         | Fence 1\n"
        "          v                 v\n"
        "GPU       Frame 0を実行      Frame 1を実行",
        GREEN,
        "FAFBFC",
        7.5,
    )
    add_spacer(doc, 4)
    add_table(
        doc,
        ["管理対象", "実装"],
        [
            ("フレーム寿命", "FrameContext / CommandAllocator / Upload Arena / deferred release / Fence"),
            ("SRV寿命", "retired descriptorをFence完了後に再利用"),
            ("状態遷移", "DirectXCommonで現在状態を追跡し、ResourceBarrierを発行"),
            ("追加バリア", "UAV barrier / aliasing barrierを共通処理として用意"),
        ],
        [2600, 7200],
        PALE_GREEN,
        8.0,
    )
    add_callout(
        doc,
        "実装方針",
        "リソースの「いつ再利用できるか」と「今どの状態で使えるか」を分けて追跡します。寿命管理と状態遷移を別々に見ることで、GPU参照中の再利用や誤った状態からの参照を検出しやすくしました。",
        PALE_ORANGE,
        ORANGE,
        font_size=8.15,
    )
    add_heading(doc, "Shadow Mapとオフスクリーン描画", 2, GREEN)
    add_para(
        doc,
        "Shadow Mapではライト視点の深度マップを作成し、通常描画のPixel ShaderでSampleCmpLevelZeroによる比較を行います。開始時はDEPTH_WRITE、終了時はPIXEL_SHADER_RESOURCEへ遷移します。GPU時間はまだ定量比較していないため、資料では実装済みの描画経路と今後の計測対象として扱います。",
        after=4,
    )
    add_flow(doc, ["PIXEL_SHADER_RESOURCE", "DEPTH_WRITE\nShadow Pass", "PIXEL_SHADER_RESOURCE", "通常描画で参照"], PALE_ORANGE, ORANGE, 7.8)

    # Page 4
    page_break(doc)
    add_kicker(doc, "03  COLLISION / PERFORMANCE", ORANGE)
    add_heading(doc, "OBBと固定グリッドによる衝突判定", 1, NAVY)
    add_para(
        doc,
        "通常弾、周囲弾、爆発弾、プレイヤー、敵の詳細判定にはOBBを使用しています。AABBだけで判定すると、回転したモデルでは見た目より判定範囲が広がるためです。ただし、全ての弾と敵へ毎回OBB判定を行うと候補数が増えるため、先に固定グリッドで周辺候補を絞ります。",
        after=4,
    )
    add_codebox(
        doc,
        "固定グリッドの概念図",
        "上から見たXZ平面\n"
        "+----+----+----+----+\n"
        "|    | E  |    |    |\n"
        "+----+----+----+----+\n"
        "|    | *  |    |    |   *: 弾またはプレイヤー\n"
        "+----+----+----+----+   E: 候補になる敵\n"
        "|    | E  | E  |    |\n"
        "+----+----+----+----+\n"
        "周辺セルの敵だけを候補にし、その後でOBBの詳細判定を行う。",
        ORANGE,
        "FAFBFC",
        7.3,
    )
    add_spacer(doc, 3)
    add_para(
        doc,
        "本作の戦闘は主にXZ平面上で行われ、敵と弾が毎フレーム移動します。そのため、階層の分割・統合が必要な四分木よりも、登録と近傍検索が単純な固定グリッドの方が今回のゲームに合っていると判断しました。",
        after=4,
    )
    add_table(
        doc,
        ["方式", "特徴", "今回の判断"],
        [
            ("総当たり", "実装は単純だが、敵数・弾数が増えるほど候補数が増える", "比較用として残す"),
            ("固定グリッド", "移動物体の再登録と周辺検索が単純", "XZ平面中心の戦闘に採用"),
            ("四分木", "密度差の大きい空間を階層化できる", "毎フレーム動く敵が多く、分割管理が増える"),
            ("BVH", "静的形状群の広域判定に向く", "敵群の更新が毎フレーム必要になる"),
        ],
        [1700, 4700, 3400],
        PALE_ORANGE,
        7.55,
    )
    add_heading(doc, "空間分割の実測結果", 2, ORANGE)
    add_table(
        doc,
        ["敵数", "総当たり 中央値", "固定グリッド 中央値", "CPU時間削減"],
        [
            ("27体", "0.1056 ms", "0.0542 ms", "48.7%"),
            ("52体", "0.3645 ms", "0.0888 ms", "75.6%"),
            ("84体", "0.9304 ms", "0.1654 ms", "82.2%"),
        ],
        [1500, 2800, 3100, 2400],
        PALE_ORANGE,
        8.0,
    )
    add_table(
        doc,
        ["条件", "総当たり相当", "固定グリッド", "削減率"],
        [("84体時の候補数中央値", "7140件", "835件", "88.3%")],
        [3200, 2200, 2200, 2200],
        PALE_BLUE,
        8.0,
    )
    add_callout(
        doc,
        "計測環境",
        "CPU: 11th Gen Intel Core i7-11800H（8コア／16スレッド）、GPU: NVIDIA GeForce RTX 3060 Laptop GPU、Release x64、固定乱数シード20260708、各条件1200フレーム。数値は衝突処理単体のCPU時間であり、FPSやGPU時間の改善値ではありません。",
        PALE_GRAY,
        BLUE,
        font_size=7.85,
    )

    # Page 5
    page_break(doc)
    add_kicker(doc, "04  DATA / DEBUG / VERIFICATION", BLUE)
    add_heading(doc, "CSVとデバッグ表示で調整を早くする", 1, NAVY)
    add_para(
        doc,
        "調整のたびにC++を修正して再ビルドする手間を減らすため、プレイヤー能力、武器強化、敵設定、スポーン、レベルアップ候補、UI配置、デバッグ調整値をCSVへ分けています。",
        after=4,
    )
    add_codebox(
        doc,
        "weaponUpgradeSettings.csv の掲載例",
        "weapon,level,damage,interval,bulletCount,range\n"
        "Normal,1,10,0.45,1,0\n"
        "Normal,2,14,0.40,1,0\n"
        "Orbit,1,6,0.00,2,3.2\n"
        "Lightning,1,20,1.20,1,5.0",
        BLUE,
        "FAFBFC",
        7.7,
    )
    add_spacer(doc, 3)
    add_callout(
        doc,
        "運用",
        "CSVで調整値を外へ出し、Dear ImGuiで実行中の状態を確認します。値の変更と原因確認をC++の修正から切り離すことで、バランス調整と不具合調査を短い周期で回せるようにしました。",
        PALE_BLUE,
        BLUE,
        font_size=8.0,
    )
    add_para(
        doc,
        "CSVでは、文字列を数値へ変換する際に、文字列全体を消費したか、有限値か、許容範囲内かを確認します。Dear ImGuiでは敵数、弾数、FPS、SRV使用数、Shadow Pass統計、衝突形状、武器レベルを確認できるようにしました。",
        after=5,
    )
    add_heading(doc, "検証", 1, GREEN)
    add_para(
        doc,
        "検証で伝えたいことは、単にテストを作ったことではありません。仕様変更や調整をした後でも、CSVの読み込み、HPの範囲制限、乱数シード、シーン遷移など、壊れやすい部分を同じ条件で再確認できるようにした点です。",
        after=4,
    )
    add_table(
        doc,
        ["確認方法", "対象", "結果・用途"],
        [
            ("CPU回帰テスト", "セルキー、候補削減率、HP、CSV、武器間隔、乱数、SoundHandle", "Debug / Release PASS"),
            ("Scene Stress", "Title -> Play -> Resultを3周", "各シーン3回訪問"),
            ("SRV telemetry", "使用数 / High-watermark", "105 / 105"),
            ("衝突計測", "Spatial Grid / Brute Force切り替え", "CPU時間と候補数をCSV化"),
        ],
        [2200, 5200, 2400],
        PALE_GREEN,
        7.7,
    )

    # Page 6
    page_break(doc)
    add_kicker(doc, "05  REVIEW / NEXT", ORANGE)
    add_heading(doc, "制作中に見直した点", 1, NAVY)
    add_para(
        doc,
        "最初は機能追加を優先していましたが、PlaySceneやPlayerに処理が集まり始めると、1つの修正で確認する範囲が広がりました。そこで、変更理由ごとにクラスを分けるリファクタリングを行い、同時に負荷が増えやすい衝突判定と、再発確認が必要な処理を整理しました。",
        after=4,
    )
    add_table(
        doc,
        ["分類", "課題", "対応", "結果"],
        [
            ("リファクタリング", "シーン生成を共通基盤へ固定したくない", "GameSceneFactoryへ分離", "ゲーム側で構成を差し替え可能"),
            ("リファクタリング", "進行処理がPlaySceneへ集中", "GameplayFlowControllerへ分離", "状態遷移を一か所で追跡"),
            ("リファクタリング", "結果の受け渡しが複雑", "GameSessionへ集約", "各シーンの役割を限定"),
            ("リファクタリング", "武器処理がPlayerへ増えやすい", "PlayerWeaponControllerへ分離", "武器更新と強化を集約"),
            ("負荷対策", "全敵への詳細判定を避けたい", "固定グリッド後にOBB判定", "近傍候補へ対象を限定"),
            ("描画管理", "GPU参照中の再利用を防ぎたい", "FrameContext / Fence / 状態追跡", "再利用条件と現在状態を明示"),
            ("検証追加", "手動確認だけでは再発を拾いにくい", "CPUテスト / Scene Stress", "同じ条件で再確認可能"),
        ],
        [1700, 3200, 2500, 2400],
        PALE_GREEN,
        7.05,
    )
    add_heading(doc, "今後の改善", 1, ORANGE)
    add_para(
        doc,
        "今後は、未計測の箇所を同じ条件で比較し、実装上の仮説を数値で確認できるようにします。特にShadow MapのGPU時間と、実プレイに近い衝突ベンチマークを優先します。",
        after=3,
    )
    add_bullet(doc, "PIXまたはGPU Timestamp Queryで、Shadow Pass有効／無効、caster数、Draw Call数、ResourceBarrierの位置を同じ条件で比較する。")
    add_bullet(doc, "実プレイ中の弾数や敵配置を保存し、同じ入力を再生する形に近づけて、衝突ベンチマークを実ゲームに近い条件へ寄せる。")
    add_bullet(doc, "起動後に追加されたモデルをパス索引へ反映する更新処理、非同期ロード、ID取り違えを防ぐ型付きハンドルを検討する。")
    add_bullet(doc, "CSVスキーマとゲーム進行をGPUなしで確認できるfixtureを増やし、調整やリファクタリング後の破損を早く拾う。")
    add_callout(
        doc,
        "現在の制約",
        "Windows / DirectX 12専用です。Shadow Mapの状態遷移は実装済みですが、GPU時間や最終的なFPS維持への寄与は今後の定量確認項目です。",
        PALE_ORANGE,
        ORANGE,
        font_size=8.0,
    )

    props = doc.core_properties
    props.title = "KuboEngine プログラム説明資料"
    props.subject = "C++ / DirectX 12 自作ゲームエンジンと3Dサバイバルアクション"
    props.author = "KuboEngine Developer"
    props.keywords = "C++, DirectX 12, HLSL, Game Engine, Portfolio"
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build()
